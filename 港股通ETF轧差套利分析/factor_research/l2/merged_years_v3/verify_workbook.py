from pathlib import Path
import json,math,os
from openpyxl import load_workbook
R=Path(os.environ['ETF_REPORT_DATA_DIR']) if 'ETF_REPORT_DATA_DIR' in os.environ else Path(__file__).resolve().parent/'outputs';O=R/'01a0915e-f9e1-76e3-bbaa-2329e4b4b7e8';d=json.loads((R/'data.json').read_text());f=O/'每日择优_2025与2026扩充回测.xlsx'
w=load_workbook(f,data_only=True);form=load_workbook(f,data_only=False)
def eq(a,b):
 if b is None or b=='':assert a is None or a=='',(a,b)
 elif isinstance(b,(float,int)) and not isinstance(b,bool):assert a is not None and math.isclose(a,b,rel_tol=1e-9,abs_tol=1e-8),(a,b)
 else:assert a==b,(a,b)
for sh,key in [('每日择优','main'),('最终汇率对照','final'),('事后重选对照','reranked')]:
 s=w[sh];rows=d[key];assert len(rows)==353
 assert form[sh].max_row==361 and len(form[sh].tables)==1
 for i,r in enumerate(rows,9):
  eq(s.cell(i,1).value.strftime('%Y-%m-%d'),r['date'])
  for col,k in [(2,'phase'),(3,'symbol'),(4,'name'),(5,'score'),(6,'net_shares'),(7,'unit'),(8,'cap_shares'),(9,'net_U'),(10,'cap_U'),(11,'result'),(12,'decision'),(13,'cap_status'),(14,'reason'),(15,'data_quality_status')]:eq(s.cell(i,col).value,r.get(k))
  assert form[sh].cell(i,9).data_type=='f' and form[sh].cell(i,10).data_type=='f'
specs=[s for s in d['summary'] if s['fx']=='lag' and s['rule']!='事后剔除候选再选第一']
for i,s in enumerate(specs,9):
 expected=[s['cohort'],s['rule'],s['days'],s['selected'],s['positive'],s['zero'],s['negative'],s['accuracy'],s['error_rate'],*s['error_wilson95'],s['unknown_cap']]
 for c,v in enumerate(expected,1):eq(w['结果汇总'].cell(i,c).value,v)
clean=[s for s in d['quality_summary'] if s['cohort'] in ['2025跨期回放','2026新增29日','2026后续检验73日']]
for i,s in enumerate(clean,25):
 for c,v in enumerate([s['cohort'],s['rule'],s['days'],s['selected'],s['positive'],s['zero'],s['negative'],s['accuracy'],s['error_rate']],1):eq(w['结果汇总'].cell(i,c).value,v)
for i,s in enumerate([s for s in d['monthly'] if s['rule']=='剔除原选中满额样本'],35):
 for c,v in enumerate([s['month'],s['days'],s['selected'],s['positive'],s['zero'],s['negative'],s['accuracy'],s['error_rate'],s['unknown_cap']],1):eq(w['结果汇总'].cell(i,c).value,v)
errors=[(s.title,c.coordinate,c.value) for s in w for row in s for c in row if c.data_type=='e'];assert not errors,errors
out=dict(passed=True,detail_rows_checked=1059,summary_rows_checked=len(specs)+len(clean),all_exported_formula_caches_match_independent_python=True,excel_error_cells=errors)
(O/'verification.json').write_text(json.dumps(out,ensure_ascii=False,indent=2));print(out)
